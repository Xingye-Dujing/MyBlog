"""DOCX renderer that fills the bundled template without losing its styles."""
from __future__ import annotations

import math
import os
import re
from typing import Any, Dict, List, Optional, Tuple

from docx import Document as load_document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .docx_utils import (
    set_paragraph_bottom_border,
    replace_in_container,
    enable_update_fields,
)
from .errors import TemplateError
from .images import image_size_cm, resolve_docx_width_cm, align_value
from .metadata import apply_docx_metadata, placeholder_mapping
from .models import Block, Document, Inline, plain_text
from .template import Template

_CAPTION_NUMBER_RE = re.compile(r"^\s*(?:图|圖|表)\s*\d+\s*[：:、.\-—\s]*")

_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _set_east_asia(run, name: str) -> None:
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), run.font.name or "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), run.font.name or "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), "Times New Roman")


class Inserter:
    def __init__(self, parent, index: int):
        self.parent = parent
        self.index = index

    def add(self, element) -> None:
        self.parent.insert(self.index, element)
        self.index += 1


class DocxRenderer:
    def __init__(self, document: Document, template: Template):
        self.document = document
        self.template = template
        self.config = template.config
        self.style_map: Dict[str, str] = dict(self.config.get("styles") or {})
        self.toc = self.config.get("toc") or {}
        self.table_no = 0
        self.figure_no = 0

    def style(self, key: str, fallback: str = "Normal"):
        name = self.style_map.get(key) or fallback
        try:
            return self.doc.styles[name]
        except KeyError:
            raise TemplateError("模板缺少样式：%s（映射键 %s）" % (name, key))

    def run(self, metadata: Dict[str, Any], output: Optional[str] = None, keep_anchor: bool = False):
        self.metadata = metadata
        self.keep_anchor = keep_anchor
        self.doc = load_document(str(self.template.path_for("template.docx")))
        self._apply_metadata()
        toc_anchor = self._find_anchor("{{TOC}}")
        content_anchor = self._find_anchor("{{CONTENT}}")
        toc_enabled, entries = self._toc_entries()
        if toc_enabled:
            self._render_toc(toc_anchor, entries)
        else:
            self._remove_toc_page(toc_anchor)
            self._remove_section_break()
        self._render_body(content_anchor, False)
        self._apply_body_placeholders()
        if toc_enabled and self.toc.get("field", True) and entries:
            enable_update_fields(self.doc)
        if output:
            self.doc.save(output)
        return self.doc

    def _apply_metadata(self):
        apply_docx_metadata(self.doc, self.metadata, self.config)

    def _find_anchor(self, text: str):
        for paragraph in self.doc.paragraphs:
            if paragraph.text.strip() == text:
                return paragraph
        raise TemplateError("模板中未找到内容锚点 %s" % text)

    def _toc_entries(self) -> Tuple[bool, List[Tuple[str, int]]]:
        if self.toc.get("enabled", True) is False or self.metadata.get("toc") is False:
            return False, []
        levels = int(self.toc.get("levels") or 2)
        pages = estimate_pages(self.document)
        entries: List[Tuple[str, int]] = []
        for block in self.document.blocks:
            if block.type == "heading" and block.level <= levels:
                text = plain_text(block.inlines).strip()
                if text:
                    entries.append((text, pages.get(id(block), 1)))
            elif block.type == "bibliography":
                entries.append(("参考文献", pages.get(id(block), entries[-1][1] if entries else 1)))
        return bool(entries), entries

    def _render_toc(self, anchor, entries: List[Tuple[str, int]]) -> None:
        if not entries:
            self._remove_toc_page(anchor)
            return
        section = self.doc.sections[0]
        available_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
        parent = anchor._p.getparent()
        index = list(parent).index(anchor._p)
        inserter = Inserter(parent, index)
        paragraphs = []
        use_field = bool(self.toc.get("field", True))
        levels = int(self.toc.get("levels") or 2)
        for entry_index, (text, page) in enumerate(entries):
            paragraph = self._new_paragraph("toc_entry")
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Cm(max(2.0, available_cm - 0.01)), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
            )
            if entry_index == 0 and use_field:
                _toc_field_begin(paragraph, levels)
            run = paragraph.add_run(text)
            _set_east_asia(run, "宋体")
            run.add_tab()
            page_run = paragraph.add_run(str(page))
            _set_east_asia(page_run, "宋体")
            paragraphs.append(paragraph)
            inserter.add(paragraph._p)
        if use_field:
            _toc_field_end(paragraphs[-1])
        parent.remove(anchor._p)

    def _remove_section_break(self) -> None:
        for paragraph in list(self.doc.paragraphs):
            p_pr = paragraph._p.pPr
            if p_pr is not None and p_pr.find(qn("w:sectPr")) is not None:
                parent = paragraph._p.getparent()
                if parent is not None:
                    parent.remove(paragraph._p)
                break

    def _remove_toc_page(self, anchor) -> None:
        title_style = self.style_map.get("toc_title")
        for paragraph in list(self.doc.paragraphs):
            if title_style and paragraph.style is not None and paragraph.style.name == title_style:
                parent = paragraph._p.getparent()
                if parent is not None:
                    parent.remove(paragraph._p)
                break
        parent = anchor._p.getparent()
        if parent is not None:
            parent.remove(anchor._p)

    def _new_paragraph(self, style_key: str, fallback: str = "Normal"):
        paragraph = self.doc.add_paragraph()
        try:
            paragraph.style = self.style(style_key, fallback)
        except TemplateError:
            raise
        parent = paragraph._p.getparent()
        parent.remove(paragraph._p)
        return paragraph

    def _new_table(self, rows: int, cols: int):
        table = self.doc.add_table(rows=rows, cols=cols)
        parent = table._tbl.getparent()
        parent.remove(table._tbl)
        return table

    def _render_body(self, content_anchor, toc_enabled: bool) -> None:
        parent = content_anchor._p.getparent()
        index = list(parent).index(content_anchor._p)
        inserter = Inserter(parent, index)
        first = True
        title = str(self.metadata.get("title") or "").strip()
        if title:
            for line in title.splitlines():
                paragraph = self._new_paragraph("title", "Title")
                if first and toc_enabled:
                    paragraph.paragraph_format.page_break_before = True
                run = paragraph.add_run(line.strip())
                _set_east_asia(run, "黑体")
                inserter.add(paragraph._p)
                first = False
        subtitle = str(self.metadata.get("subtitle") or "").strip()
        if subtitle:
            for line in subtitle.splitlines():
                paragraph = self._new_paragraph("subtitle", "Subtitle")
                if first and toc_enabled:
                    paragraph.paragraph_format.page_break_before = True
                run = paragraph.add_run(line.strip())
                _set_east_asia(run, "黑体")
                inserter.add(paragraph._p)
                first = False
        abstract = _as_text(self.metadata.get("abstract"))
        if abstract:
            paragraph = self._new_paragraph("abstract")
            if first and toc_enabled:
                paragraph.paragraph_format.page_break_before = True
            label = paragraph.add_run("摘  要：")
            label.bold = False
            _set_east_asia(label, "黑体")
            body = paragraph.add_run(abstract)
            _set_east_asia(body, "宋体")
            inserter.add(paragraph._p)
            first = False
        keywords = self.metadata.get("keywords")
        if keywords:
            paragraph = self._new_paragraph("keywords")
            if first and toc_enabled:
                paragraph.paragraph_format.page_break_before = True
            label = paragraph.add_run("关键词：")
            label.bold = False
            _set_east_asia(label, "黑体")
            body = paragraph.add_run(_as_text(keywords))
            _set_east_asia(body, "宋体")
            inserter.add(paragraph._p)
            first = False
            spacer = self._new_paragraph("paragraph")
            spacer.paragraph_format.first_line_indent = Pt(0)
            inserter.add(spacer._p)

        for block in self.document.blocks:
            self._render_block(block, inserter)
        if not self.keep_anchor:
            parent.remove(content_anchor._p)
        if first:
            paragraph = self._new_paragraph("paragraph")
            if toc_enabled:
                paragraph.paragraph_format.page_break_before = True
            inserter.add(paragraph._p)

    def _render_block(self, block: Block, inserter: Inserter) -> None:
        if block.type == "heading":
            level = min(4, max(1, block.level))
            paragraph = self._new_paragraph("h%d" % level, "Heading %d" % level)
            self._add_inlines(paragraph, block.inlines)
            inserter.add(paragraph._p)
        elif block.type == "paragraph":
            paragraph = self._new_paragraph("paragraph")
            self._add_inlines(paragraph, block.inlines)
            inserter.add(paragraph._p)
        elif block.type == "list":
            self._render_list(block, inserter)
        elif block.type == "code_block":
            paragraph = self._new_paragraph("code_block")
            run = paragraph.add_run()
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            _set_east_asia(run, "宋体")
            lines = block.text.rstrip("\n").split("\n")
            for index, line in enumerate(lines):
                if index:
                    run.add_break()
                run.add_text(line)
            inserter.add(paragraph._p)
        elif block.type == "quote":
            paragraph = self._new_paragraph("quote")
            self._add_inlines(paragraph, block.inlines)
            inserter.add(paragraph._p)
        elif block.type == "table":
            self._render_table(block, inserter)
        elif block.type == "image":
            self._render_image(block, inserter)
        elif block.type == "latex_figure":
            self._render_latex_figure(block, inserter)
        elif block.type == "pagebreak":
            paragraph = self._new_paragraph("paragraph")
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            inserter.add(paragraph._p)
        elif block.type == "hr":
            paragraph = self._new_paragraph("paragraph")
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            set_paragraph_bottom_border(paragraph)
            inserter.add(paragraph._p)
        elif block.type == "math_block":
            paragraph = self._new_paragraph("math")
            run = paragraph.add_run(block.text)
            run.font.name = "Cambria Math"
            run.italic = True
            _set_east_asia(run, "宋体")
            inserter.add(paragraph._p)
        elif block.type == "bibliography":
            self._render_bibliography(block, inserter)

    def _render_list(self, block: Block, inserter: Inserter) -> None:
        counters: Dict[int, int] = {}
        self._render_list_level(block, inserter, counters)

    def _render_list_level(self, block: Block, inserter: Inserter, counters: Dict[int, int]) -> None:
        for item in block.items:
            level = min(3, max(1, item.level))
            key = "ordered_list_l%d" % level if block.ordered else "bullet_list_l%d" % level
            paragraph = self._new_paragraph(key)
            if block.ordered:
                if id(block) not in counters:
                    counters[id(block)] = max(0, int(block.start or 1) - 1)
                counters[id(block)] += 1
                prefix = "%d. " % counters[id(block)]
            else:
                prefix = ["•", "◦", "▪"][level - 1] + " "
            prefix_run = paragraph.add_run(prefix)
            _set_east_asia(prefix_run, "宋体")
            self._add_inlines(paragraph, item.inlines)
            inserter.add(paragraph._p)
            for child in item.children:
                if child.type == "list":
                    self._render_list_level(child, inserter, counters)

    def _render_table(self, block: Block, inserter: Inserter) -> None:
        rows = block.rows or []
        if not rows:
            return
        cols = max(len(row) for row in rows)
        if cols == 0:
            return
        for row in rows:
            while len(row) < cols:
                row.append([])
        caption = _strip_caption_number(block.caption or "")
        if caption:
            self.table_no += 1
            paragraph = self._new_paragraph("caption")
            paragraph.paragraph_format.keep_with_next = True
            label = "表 %d　" % self.table_no
            run = paragraph.add_run(label)
            run.bold = True
            _set_east_asia(run, "黑体")
            run2 = paragraph.add_run(caption)
            _set_east_asia(run2, "宋体")
            inserter.add(paragraph._p)
        table = self._new_table(len(rows), cols)
        try:
            table.style = self.style("table", "Table Grid")
        except KeyError:
            table.style = self.doc.styles["Table Grid"]
        section = self.doc.sections[0]
        available_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
        table.autofit = False
        column_width = available_cm / cols
        for row_index, row in enumerate(rows):
            for col_index in range(cols):
                cell = table.cell(row_index, col_index)
                cell.width = Cm(column_width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                try:
                    paragraph.style = self.style("table_text", "Normal")
                except KeyError:
                    pass
                align = block.aligns[col_index] if col_index < len(block.aligns) else ""
                if row_index == 0:
                    paragraph.alignment = _ALIGN_MAP.get(align, WD_ALIGN_PARAGRAPH.CENTER)
                else:
                    paragraph.alignment = _ALIGN_MAP.get(align, WD_ALIGN_PARAGRAPH.LEFT)
                self._add_inlines(paragraph, row[col_index], force_bold=(row_index == 0))
        _three_line_table(table, header_rows=1)
        inserter.add(table._tbl)

    def _render_image(self, block: Block, inserter: Inserter) -> None:
        image = block.image or {}
        path = str(image.get("path") or "")
        if not path:
            return
        if re.match(r"^[a-z]+://", path, re.I):
            return
        full_path = path
        if not os.path.isabs(full_path):
            full_path = os.path.normpath(os.path.join(self.document.base_dir or os.getcwd(), path))
        if not os.path.isfile(full_path):
            self.document.add_issue("error", "图片不存在：%s" % path, block.line, "image-missing")
            return
        section = self.doc.sections[0]
        available = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
        max_width = float((self.config.get("image") or {}).get("max_width_cm") or available)
        try:
            width_cm = resolve_docx_width_cm(str(image.get("width") or ""), full_path, available, max_width)
            natural_w, natural_h = image_size_cm(full_path)
            if natural_w > 0 and natural_h > 0:
                max_height = 20.0
                if natural_h / natural_w * width_cm > max_height:
                    width_cm = max_height * natural_w / natural_h
        except Exception:
            width_cm = min(available, max_width)
        paragraph = self._new_paragraph("image")
        paragraph.alignment = _ALIGN_MAP.get(align_value(str(image.get("align") or "")), WD_ALIGN_PARAGRAPH.CENTER)
        run = paragraph.add_run()
        try:
            run.add_picture(full_path, width=Cm(width_cm))
        except Exception as exc:
            self.document.add_issue("error", "图片插入失败：%s (%s)" % (path, exc), block.line, "image-render")
            return
        inserter.add(paragraph._p)
        caption = _strip_caption_number(block.caption or image.get("caption") or image.get("alt") or "")
        if caption:
            self.figure_no += 1
            caption_paragraph = self._new_paragraph("caption")
            label = caption_paragraph.add_run("图 %d　" % self.figure_no)
            label.bold = True
            _set_east_asia(label, "黑体")
            body = caption_paragraph.add_run(caption)
            _set_east_asia(body, "宋体")
            inserter.add(caption_paragraph._p)

    def _render_latex_figure(self, block: Block, inserter: Inserter) -> None:
        image = block.image or {}
        path = str(image.get("path") or "")
        if path and not re.match(r"^[a-z]+://", path, re.I) and not os.path.isabs(path):
            full = os.path.normpath(os.path.join(self.document.base_dir or os.getcwd(), path))
            if os.path.isfile(full):
                self._render_image(block, inserter)
                return
        self.figure_no += 1
        caption = _strip_caption_number(block.caption or "")
        paragraph = self._new_paragraph("caption")
        label = paragraph.add_run("图 %d　" % self.figure_no)
        label.bold = True
        _set_east_asia(label, "黑体")
        body = paragraph.add_run(caption or "LaTeX 图表")
        _set_east_asia(body, "宋体")
        note = paragraph.add_run("（DOCX 占位；TeX 输出使用 LaTeX 绘制）")
        note.font.size = Pt(8)
        _set_east_asia(note, "宋体")
        inserter.add(paragraph._p)
        self.document.add_issue(
            "warning",
            "LaTeX 图表在 DOCX 中使用占位，请提供 fallback 图片以获得可视图表：%s" % (caption or "未命名"),
            block.line,
            "latex-figure-docx",
        )

    def _render_bibliography(self, block: Block, inserter: Inserter) -> None:
        paragraph = self._new_paragraph("bibliography_title")
        run = paragraph.add_run("参考文献")
        run.bold = True
        _set_east_asia(run, "黑体")
        inserter.add(paragraph._p)
        for inlines in block.items:
            item = self._new_paragraph("bibliography")
            self._add_inlines(item, inlines)
            inserter.add(item._p)

    def _add_inlines(
        self,
        paragraph,
        inlines: List[Inline],
        bold: bool = False,
        italic: bool = False,
        strike: bool = False,
        force_bold: bool = False,
    ) -> None:
        for item in inlines:
            if item.type == "text":
                self._add_text_run(paragraph, item.text, bold or force_bold, italic, strike)
            elif item.type == "strong":
                self._add_inlines(paragraph, item.children, True, italic, strike, force_bold)
            elif item.type == "em":
                self._add_inlines(paragraph, item.children, bold, True, strike, force_bold)
            elif item.type == "strike":
                self._add_inlines(paragraph, item.children, bold, italic, True, force_bold)
            elif item.type == "code":
                run = paragraph.add_run(item.text)
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                _set_east_asia(run, "宋体")
                if bold or force_bold:
                    run.bold = True
                if italic:
                    run.italic = True
                if strike:
                    run.font.strike = True
            elif item.type == "link":
                text = plain_text(item.children) or item.url
                _add_hyperlink(paragraph, item.url, text)

            elif item.type == "math":
                run = paragraph.add_run(item.text)
                run.font.name = "Cambria Math"
                run.italic = True
                _set_east_asia(run, "宋体")
                if bold or force_bold:
                    run.bold = True
            elif item.type == "break":
                paragraph.add_run().add_break()
            else:
                self._add_inlines(paragraph, item.children, bold, italic, strike, force_bold)

    def _add_text_run(self, paragraph, text: str, bold: bool, italic: bool, strike: bool) -> None:
        if not text:
            return
        run = paragraph.add_run()
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if strike:
            run.font.strike = True
        for index, part in enumerate(str(text).split("\n")):
            if index:
                run.add_break()
            run.add_text(part)

    def _apply_body_placeholders(self) -> None:
        mapping = placeholder_mapping(self.config, self.metadata)
        try:
            replace_in_container(self.doc, mapping)
        except Exception:
            pass


def _strip_caption_number(text: str) -> str:
    return _CAPTION_NUMBER_RE.sub("", str(text or "")).strip()


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "；".join(str(item) for item in value if str(item).strip())
    return str(value)


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    try:
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    except Exception:
        paragraph.add_run(text)
        return
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "宋体")
    r_pr.append(r_fonts)
    run.append(r_pr)
    for index, part in enumerate(str(text).split("\n")):
        if index:
            run.append(OxmlElement("w:br"))
        if part:
            text_el = OxmlElement("w:t")
            text_el.set(qn("xml:space"), "preserve")
            text_el.text = part
            run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _toc_field_begin(paragraph, levels: int = 2) -> None:
    p = paragraph._p
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-%d" \\h \\z \\u ' % max(1, int(levels))
    instr_run.append(instr)
    sep_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    sep_run.append(separate)
    insert_at = 1 if p.pPr is not None else 0
    for offset, run in enumerate((begin_run, instr_run, sep_run)):
        p.insert(insert_at + offset, run)


def _toc_field_end(paragraph) -> None:
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _three_line_table(table, header_rows: int = 1) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        element = OxmlElement("w:%s" % edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "12")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    for edge in ("left", "right", "insideH", "insideV"):
        element = OxmlElement("w:%s" % edge)
        element.set(qn("w:val"), "nil")
        borders.append(element)
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    following = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
    insert_at = len(tbl_pr)
    for index, child in enumerate(tbl_pr):
        tag = child.tag.split("}")[-1]
        if ("w:" + tag) in following:
            insert_at = index
            break
    tbl_pr.insert(insert_at, borders)
    for row_index in range(min(header_rows, len(table.rows))):
        row = table.rows[row_index]
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(tbl_header)
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            cell_borders = OxmlElement("w:tcBorders")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "0")
            bottom.set(qn("w:color"), "000000")
            cell_borders.append(bottom)
            v_align = tc_pr.find(qn("w:vAlign"))
            if v_align is not None:
                tc_pr.insert(list(tc_pr).index(v_align), cell_borders)
            else:
                tc_pr.append(cell_borders)


def estimate_pages(document: Document) -> Dict[int, int]:
    capacity = 43
    chars_per_line = 38
    page = 1
    lines = 0
    abstract = _as_text(document.metadata.get("abstract"))
    keywords = _as_text(document.metadata.get("keywords"))
    lines += 2
    if abstract:
        lines += max(1, math.ceil(len(abstract) / chars_per_line))
    if keywords:
        lines += max(1, math.ceil(len(keywords) / chars_per_line))
    result: Dict[int, int] = {}

    def consume(count: int):
        nonlocal page, lines
        lines += max(1, count)
        while lines > capacity:
            page += 1
            lines -= capacity

    for block in document.blocks:
        if block.type == "heading":
            if lines > capacity - 3:
                page += 1
                lines = 0
            result[id(block)] = page
            text = plain_text(block.inlines)
            consume(max(1, math.ceil(len(text) / chars_per_line)))
        elif block.type == "paragraph":
            consume(max(1, math.ceil(len(plain_text(block.inlines)) / chars_per_line)))
        elif block.type == "list":
            _consume_list_lines(block, consume, chars_per_line)
        elif block.type == "code_block":
            consume(max(1, len(block.text.splitlines())))
        elif block.type == "table":
            consume(max(2, len(block.rows) + 1))
        elif block.type == "image":
            consume(12)
        elif block.type == "math_block":
            consume(max(1, len(block.text.splitlines())))
        elif block.type == "pagebreak":
            page += 1
            lines = 0
        elif block.type == "bibliography":
            result[id(block)] = page
            consume(max(2, len(block.items) + 1))
    return result


def _consume_list_lines(block: Block, consume, chars_per_line: int) -> None:
    for item in block.items:
        consume(max(1, math.ceil(len(plain_text(item.inlines)) / chars_per_line)))
        for child in item.children:
            if child.type == "list":
                _consume_list_lines(child, consume, chars_per_line)


def render_docx(document: Document, template: Template, output: Optional[str] = None, keep_anchor: bool = False):
    return DocxRenderer(document, template).run(document.metadata, output, keep_anchor=keep_anchor)
