"""Metadata application for DOCX and TeX."""
from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Any, Dict, List, Optional

from docx import Document as load_docx

from .docx_utils import replace_in_container
from .errors import TemplateError
from .escape import escape


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "；".join(str(v) for v in value if str(v).strip())
    if isinstance(value, dict):
        return "；".join("%s: %s" % (k, v) for k, v in value.items())
    return str(value)


def _parse_datetime(value: Any) -> Optional[datetime]:
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except Exception:
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(str(value), fmt)
            except Exception:
                continue
    return None


def core_mapping(config: Dict[str, Any]) -> Dict[str, str]:
    meta = config.get("metadata") or {}
    mapping = dict(meta.get("core_properties") or {})
    if not mapping:
        mapping = {
            "title": "title",
            "author": "author",
            "subject": "subject",
            "keywords": "keywords",
            "version": "comments",
            "department": "category",
        }
    return mapping


def placeholder_mapping(
    config: Dict[str, Any], metadata: Dict[str, Any], include_missing: bool = True
) -> Dict[str, str]:
    defaults = {
        "{{TITLE}}": "title",
        "{{AUTHOR}}": "author",
        "{{DATE}}": "date",
        "{{DEPT}}": "department",
        "{{SUBJECT}}": "subject",
        "{{VERSION}}": "version",
    }
    mapping: Dict[str, str] = {}
    for placeholder, field in defaults.items():
        if not include_missing and field not in metadata:
            continue
        mapping[placeholder] = _text(metadata.get(field))
    meta = config.get("metadata") or {}
    for group in ("header", "footer", "body"):
        for item in meta.get(group) or []:
            if not isinstance(item, dict) or not item.get("placeholder"):
                continue
            field = item.get("field") or ""
            if not include_missing and field not in metadata:
                continue
            mapping[str(item["placeholder"])] = _text(metadata.get(field))
    return mapping


def apply_docx_metadata(
    doc, metadata: Dict[str, Any], config: Dict[str, Any], clear_missing: bool = True
) -> None:
    mapping = core_mapping(config)
    props = doc.core_properties
    for field, prop in mapping.items():
        if field not in metadata:
            continue
        if field in ("created", "modified"):
            parsed = _parse_datetime(metadata.get(field))
            if parsed is not None:
                try:
                    setattr(props, prop, parsed)
                except Exception:
                    pass
            continue
        if field == "revision":
            try:
                setattr(props, prop, int(metadata.get(field) or 1))
            except Exception:
                pass
            continue
        value = _text(metadata.get(field))
        if hasattr(props, prop):
            try:
                setattr(props, prop, value)
            except Exception:
                pass
    if metadata.get("title") and not props.title:
        props.title = _text(metadata.get("title"))
    if metadata.get("author") and not props.author:
        props.author = _text(metadata.get("author"))

    replacement = placeholder_mapping(config, metadata, include_missing=clear_missing)
    seen = set()
    for section in doc.sections:
        for container in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            key = id(container._element)
            if key in seen:
                continue
            seen.add(key)
            try:
                replace_in_container(container, replacement)
            except Exception:
                continue


def apply_tex_metadata_patterns(tex: str, updates: Dict[str, Any], config: Dict[str, Any]) -> str:
    meta = config.get("metadata") or {}
    patterns = meta.get("meta_patterns") or config.get("meta_patterns") or {}
    if not patterns:
        patterns = {
            "title": r"\\title\{([^{}]*)\}",
            "author": r"\\author\{([^{}]*)\}",
            "date": r"\\date\{([^{}]*)\}",
        }
    result = tex
    for field, pattern in patterns.items():
        if field not in updates:
            continue
        value = escape(_text(updates.get(field)))
        try:
            result = re.sub(pattern, lambda m, v=value: m.group(0).replace(m.group(1), v) if m.groups() else m.group(0), result)
        except re.error as exc:
            raise TemplateError("meta_patterns 正则无效（%s）：%s" % (field, exc))
    return result


def edit_metadata(
    path: str,
    updates: Dict[str, Any],
    header_placeholders: Optional[Dict[str, str]] = None,
    output: Optional[str] = None,
    template_config: Optional[Dict[str, Any]] = None,
) -> str:
    config = template_config or {}
    lower = str(path).lower()
    if lower.endswith(".docx"):
        doc = load_docx(path)
        seed: Dict[str, Any] = {}
        for field, prop in core_mapping(config).items():
            if hasattr(doc.core_properties, prop):
                value = getattr(doc.core_properties, prop)
                if value not in (None, ""):
                    seed[field] = value
        seed.update(updates)
        apply_docx_metadata(doc, seed, config, clear_missing=False)
        if header_placeholders:
            seen = set()
            for section in doc.sections:
                for container in (
                    section.header,
                    section.footer,
                    section.first_page_header,
                    section.first_page_footer,
                    section.even_page_header,
                    section.even_page_footer,
                ):
                    key = id(container._element)
                    if key in seen:
                        continue
                    seen.add(key)
                    replace_in_container(container, header_placeholders)
        target = output or path
        doc.save(target)
        return target
    if lower.endswith(".tex"):
        with open(path, "r", encoding="utf-8") as handle:
            source = handle.read()
        result = apply_tex_metadata_patterns(source, updates, config)
        target = output or path
        with open(target, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(result)
        return target
    raise TemplateError("仅支持 .docx 或 .tex 文件")
