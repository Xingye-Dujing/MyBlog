from pathlib import Path

from docx import Document as open_docx
from docx.oxml.ns import qn
from docx.shared import Pt

from mdforge.docx_utils import replace_in_paragraph
from mdforge.metadata import apply_docx_metadata, apply_tex_metadata_patterns, edit_metadata


def test_replace_across_runs_preserves_format():
    doc = open_docx()
    paragraph = doc.add_paragraph()
    first = paragraph.add_run("{{TI")
    first.bold = True
    paragraph.add_run("TLE}}")
    paragraph.add_run(" 结尾")
    changed = replace_in_paragraph(paragraph, {"{{TITLE}}": "新标题"})
    assert changed is True
    assert paragraph.text == "新标题 结尾"
    assert any(run.bold for run in paragraph.runs if run.text == "新标题")


def test_core_properties_and_headers():
    doc = open_docx()
    doc.add_paragraph("正文")
    header = doc.sections[0].header
    paragraph = header.paragraphs[0]
    paragraph.add_run("{{TITLE}}")
    paragraph.add_run(" · ")
    paragraph.add_run("{{DEPT}}")
    config = {
        "metadata": {
            "core_properties": {
                "title": "title",
                "author": "author",
                "keywords": "keywords",
                "version": "comments",
                "department": "category",
            },
            "header": [
                {"placeholder": "{{TITLE}}", "field": "title"},
                {"placeholder": "{{DEPT}}", "field": "department"},
            ],
        }
    }
    metadata = {
        "title": "测试标题",
        "author": "测试作者",
        "keywords": ["A", "B"],
        "version": "v2",
        "department": "研发部",
    }
    apply_docx_metadata(doc, metadata, config)
    assert doc.core_properties.title == "测试标题"
    assert doc.core_properties.author == "测试作者"
    assert doc.core_properties.keywords == "A；B"
    assert doc.core_properties.comments == "v2"
    assert doc.core_properties.category == "研发部"
    assert header.paragraphs[0].text == "测试标题 · 研发部"


def test_edit_existing_docx(tmp_path):
    doc = open_docx()
    doc.core_properties.title = "旧标题"
    doc.core_properties.author = "旧作者"
    header = doc.sections[0].header
    header.paragraphs[0].add_run("{{AUTHOR}}")
    doc.add_paragraph("hello")
    source = tmp_path / "input.docx"
    doc.save(str(source))
    output = tmp_path / "output.docx"
    edit_metadata(
        str(source),
        {"title": "新标题", "author": "李四"},
        header_placeholders={"{{AUTHOR}}": "李四"},
        output=str(output),
    )
    result = open_docx(str(output))
    assert result.core_properties.title == "新标题"
    assert result.sections[0].header.paragraphs[0].text == "李四"


def test_tex_meta_patterns():
    config = {"metadata": {"meta_patterns": {"title": r"pdftitle=\{([^{}]*)\}"}}}
    source = "\\hypersetup{pdftitle={旧标题}}\n"
    assert "新标题" in apply_tex_metadata_patterns(source, {"title": "新标题"}, config)


def test_tex_meta_defaults():
    source = "\\title{旧标题}\n\\author{旧作者}\n\\date{2020-01-01}\n"
    result = apply_tex_metadata_patterns(source, {"title": "新标题", "author": "李四", "date": "2026-10-21"}, {})
    assert "\\title{新标题}" in result
    assert "\\author{李四}" in result
    assert "\\date{2026-10-21}" in result


def test_partial_metadata_update_keeps_other_properties():
    doc = open_docx()
    doc.core_properties.title = "旧标题"
    doc.core_properties.author = "旧作者"
    doc.core_properties.keywords = "旧关键词"
    config = {"metadata": {"core_properties": {"title": "title", "author": "author", "keywords": "keywords"}}}
    apply_docx_metadata(doc, {"author": "新作者"}, config)
    assert doc.core_properties.author == "新作者"
    assert doc.core_properties.title == "旧标题"
    assert doc.core_properties.keywords == "旧关键词"
