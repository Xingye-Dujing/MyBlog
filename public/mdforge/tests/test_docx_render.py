import hashlib
from pathlib import Path

from docx import Document as open_docx
from docx.oxml.ns import qn

from mdforge.docx_renderer import render_docx
from mdforge.parser import parse_markdown
from mdforge.template import load_template

ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "examples" / "sample.md"


def make_document():
    source = SAMPLE.read_text(encoding="utf-8")
    return parse_markdown(source, "sample.md", base_dir=str(ROOT / "examples"))


def test_render_reference_layout(tmp_path):
    document = make_document()
    template = load_template("docx", "academic")
    output = tmp_path / "out.docx"
    render_docx(document, template, str(output))
    assert output.is_file()
    doc = open_docx(str(output))
    assert doc.core_properties.title == "从文本到版式"
    assert doc.core_properties.author == "MdForge"
    assert len(doc.sections) == 2
    section = doc.sections[0]
    assert round(section.page_width.cm, 2) == 21.0
    assert round(section.page_height.cm, 2) == 29.7
    assert abs(section.left_margin.cm - 3.175) < 0.01
    assert abs(section.top_margin.cm - 2.54) < 0.01
    pg_num = doc.sections[1]._sectPr.find(qn("w:pgNumType"))
    assert pg_num is not None and pg_num.get(qn("w:start")) == "1"
    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert "目录" in texts
    assert "从文本到版式" in texts
    assert "摘  要：" in next(text for text in texts if text.startswith("摘  要："))
    assert "关键词：" in next(text for text in texts if text.startswith("关键词："))
    assert "引言" in texts
    assert "参考文献" in texts
    assert "{{TOC}}" not in texts
    assert "{{CONTENT}}" not in texts
    assert len(doc.tables) == 1
    assert len(doc.inline_shapes) == 1
    styles = {paragraph.style.name for paragraph in doc.paragraphs}
    assert "Heading 1" in styles
    assert "ZEN Body" in styles
    assert "ZEN Caption" in styles
    assert "ZEN Reference" in styles
    body_style = doc.styles["ZEN Body"]
    assert body_style.font.size.pt == 10.5
    assert round(body_style.paragraph_format.first_line_indent.pt, 1) == 21.0
    h1 = doc.styles["Heading 1"]
    assert h1.font.size.pt == 12.0
    assert h1.font.bold is False
    update = doc.settings.element.find(qn("w:updateFields"))
    assert update is not None and update.get(qn("w:val")) == "true"
    code = next(paragraph for paragraph in doc.paragraphs if paragraph.style.name == "ZEN Code")
    assert len(code._p.xpath(".//w:br")) >= 2


def test_keep_anchor(tmp_path):
    document = make_document()
    template = load_template("docx", "academic")
    output = tmp_path / "out-anchor.docx"
    render_docx(document, template, str(output), keep_anchor=True)
    doc = open_docx(str(output))
    assert "{{CONTENT}}" in [paragraph.text for paragraph in doc.paragraphs]


def test_toc_disabled(tmp_path):
    document = make_document()
    document.metadata["toc"] = False
    template = load_template("docx", "academic")
    output = tmp_path / "out-notoc.docx"
    render_docx(document, template, str(output))
    doc = open_docx(str(output))
    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert "目录" not in texts
    assert "从文本到版式" in texts


def test_template_is_not_modified(tmp_path):
    template = load_template("docx", "academic")
    template_path = template.path_for("template.docx")
    before = hashlib.sha256(template_path.read_bytes()).hexdigest()
    render_docx(make_document(), template, str(tmp_path / "out.docx"))
    after = hashlib.sha256(template_path.read_bytes()).hexdigest()
    assert before == after
