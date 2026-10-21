"""Build the bundled academic DOCX template from the reference layout."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .docx_utils import set_style_outline


def _set_east_asia(style_or_run, font_name: str) -> None:
    element = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    r_pr = element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), "Times New Roman")


def _paragraph_style(doc, name: str, base: str = "Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles[base]
        return style


def _configure_base_styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    _set_east_asia(normal, "宋体")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = False

    heading_specs = [
        ("Heading 1", 12.0, "黑体", False, 0, 3.15, 3.15),
        ("Heading 2", 10.5, "宋体", True, 1, 2.6, 2.6),
        ("Heading 3", 10.5, "宋体", True, 2, 2.0, 2.0),
        ("Heading 4", 10.5, "宋体", True, 3, 1.5, 1.5),
    ]
    for name, size, east, bold, outline, before, after in heading_specs:
        style = doc.styles[name]
        style.base_style = normal
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        _set_east_asia(style, east)
        pf = style.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.line_spacing = 1.5
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.first_line_indent = Pt(0)
        pf.keep_with_next = True
        pf.widow_control = False
        set_style_outline(style, outline)

    title = _paragraph_style(doc, "ZEN Title")
    title.font.name = "Times New Roman"
    title.font.size = Pt(15)
    title.font.bold = False
    _set_east_asia(title, "黑体")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.5
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(15.75)
    title.paragraph_format.keep_with_next = True

    subtitle = _paragraph_style(doc, "ZEN Subtitle")
    subtitle.base_style = title
    subtitle.font.size = Pt(15)
    subtitle.paragraph_format.space_after = Pt(15.75)

    abstract = _paragraph_style(doc, "ZEN Abstract")
    abstract.font.size = Pt(10.5)
    _set_east_asia(abstract, "宋体")
    abstract.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.first_line_indent = Pt(21)
    abstract.paragraph_format.line_spacing = 1.5
    abstract.paragraph_format.space_before = Pt(0)
    abstract.paragraph_format.space_after = Pt(0)
    abstract.paragraph_format.widow_control = False

    keywords = _paragraph_style(doc, "ZEN Keywords")
    keywords.base_style = abstract
    keywords.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    keywords.paragraph_format.space_after = Pt(6)

    body = _paragraph_style(doc, "ZEN Body")
    body.font.size = Pt(10.5)
    _set_east_asia(body, "宋体")
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body.paragraph_format.first_line_indent = Pt(21)
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.widow_control = False

    toc_title = _paragraph_style(doc, "ZEN TOC Title")
    toc_title.font.name = "Times New Roman"
    toc_title.font.size = Pt(15)
    toc_title.font.bold = False
    _set_east_asia(toc_title, "黑体")
    toc_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.line_spacing = 1.5
    toc_title.paragraph_format.space_before = Pt(15.75)
    toc_title.paragraph_format.space_after = Pt(15.75)
    toc_title.paragraph_format.keep_with_next = True

    toc_entry = _paragraph_style(doc, "ZEN TOC Entry")
    toc_entry.font.size = Pt(10.5)
    _set_east_asia(toc_entry, "宋体")
    toc_entry.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    toc_entry.paragraph_format.line_spacing = 1.5
    toc_entry.paragraph_format.space_before = Pt(0)
    toc_entry.paragraph_format.space_after = Pt(0)

    for name in ("TOC 1", "TOC 2"):
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Times New Roman"
        style.font.size = Pt(10.5)
        _set_east_asia(style, "宋体")
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

    list_specs = [
        ("ZEN List 1", 0.74),
        ("ZEN List 2", 1.48),
        ("ZEN List 3", 2.22),
    ]
    for name, indent in list_specs:
        style = _paragraph_style(doc, name)
        style.base_style = body
        style.paragraph_format.left_indent = Cm(indent)
        style.paragraph_format.first_line_indent = Cm(-0.74)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    code = _paragraph_style(doc, "ZEN Code")
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    _set_east_asia(code, "宋体")
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.first_line_indent = Pt(0)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.left_indent = Cm(0.2)
    code.paragraph_format.right_indent = Cm(0.2)
    p_pr = code.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F3EEE3")
    p_pr.append(shd)

    quote = _paragraph_style(doc, "ZEN Quote")
    quote.base_style = body
    quote.paragraph_format.left_indent = Cm(0.74)
    quote.paragraph_format.right_indent = Cm(0.74)
    quote.paragraph_format.first_line_indent = Pt(0)
    quote.paragraph_format.space_before = Pt(3)
    quote.paragraph_format.space_after = Pt(3)

    caption = _paragraph_style(doc, "ZEN Caption")
    caption.font.size = Pt(10.5)
    _set_east_asia(caption, "宋体")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.keep_with_next = False

    image = _paragraph_style(doc, "ZEN Image")
    image.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image.paragraph_format.first_line_indent = Pt(0)
    image.paragraph_format.line_spacing = 1.0
    image.paragraph_format.space_before = Pt(6)
    image.paragraph_format.space_after = Pt(0)
    image.paragraph_format.keep_with_next = True

    table_text = _paragraph_style(doc, "ZEN Table Text")
    table_text.font.size = Pt(9)
    _set_east_asia(table_text, "宋体")
    table_text.paragraph_format.first_line_indent = Pt(0)
    table_text.paragraph_format.line_spacing = 1.0
    table_text.paragraph_format.space_before = Pt(2)
    table_text.paragraph_format.space_after = Pt(2)

    math_style = _paragraph_style(doc, "ZEN Math")
    math_style.font.name = "Cambria Math"
    math_style.font.size = Pt(10.5)
    _set_east_asia(math_style, "宋体")
    math_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_style.paragraph_format.first_line_indent = Pt(0)
    math_style.paragraph_format.space_before = Pt(6)
    math_style.paragraph_format.space_after = Pt(6)

    ref_title = _paragraph_style(doc, "ZEN Reference Title")
    ref_title.font.name = "Times New Roman"
    ref_title.font.size = Pt(10.5)
    ref_title.font.bold = True
    _set_east_asia(ref_title, "黑体")
    ref_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ref_title.paragraph_format.first_line_indent = Pt(0)
    ref_title.paragraph_format.line_spacing = 1.5
    ref_title.paragraph_format.space_before = Pt(0)
    ref_title.paragraph_format.space_after = Pt(6)
    ref_title.paragraph_format.page_break_before = True
    set_style_outline(ref_title, 0)

    ref_body = _paragraph_style(doc, "ZEN Reference")
    ref_body.base_style = body
    ref_body.paragraph_format.first_line_indent = Pt(0)
    ref_body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def build_academic_template(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.175)
    section.right_margin = Cm(3.175)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)
    section.different_first_page_header_footer = False

    _configure_base_styles(doc)

    doc.add_paragraph("目录", style="ZEN TOC Title")
    doc.add_paragraph("{{TOC}}", style="ZEN TOC Entry")
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:fmt"), "decimal")
    pg_num.set(qn("w:start"), "1")
    sect_pr = body_section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is not None:
        sect_pr.insert(list(sect_pr).index(cols), pg_num)
    else:
        sect_pr.append(pg_num)
    doc.add_paragraph("{{CONTENT}}", style="ZEN Body")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


if __name__ == "__main__":
    target = Path(__file__).resolve().parent / "templates" / "docx" / "academic" / "template.docx"
    build_academic_template(target)
    print("template written:", target)
