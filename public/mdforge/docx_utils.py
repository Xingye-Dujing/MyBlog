"""Low-level python-docx helpers."""
from __future__ import annotations

import copy
import re
from typing import Any, Dict, Iterable, List, Optional, Tuple

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_paragraph_bottom_border(paragraph, color: str = "3A3027", size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_style_outline(style, level: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def add_complex_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


def enable_update_fields(doc) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def find_paragraph_by_text(doc, text: str):
    for paragraph in doc.paragraphs:
        if text in paragraph.text:
            return paragraph
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                if text in paragraph.text:
                    return paragraph
    return None


def _runs_with_positions(paragraph) -> List[Tuple[Any, int, int]]:
    runs = []
    position = 0
    for run in paragraph.runs:
        text = run.text or ""
        runs.append((run, position, position + len(text)))
        position += len(text)
    return runs


def _find_matches(full_text: str, mapping: Dict[str, str]) -> List[Tuple[int, int, str]]:
    matches: List[Tuple[int, int, str]] = []
    for placeholder, value in mapping.items():
        if not placeholder:
            continue
        start = 0
        while True:
            idx = full_text.find(placeholder, start)
            if idx < 0:
                break
            matches.append((idx, idx + len(placeholder), "" if value is None else str(value)))
            start = idx + len(placeholder)
    matches.sort(key=lambda item: (item[0], -(item[1] - item[0])))
    filtered: List[Tuple[int, int, str]] = []
    last_end = -1
    for start, end, value in matches:
        if start >= last_end:
            filtered.append((start, end, value))
            last_end = end
    return filtered


def replace_in_paragraph(paragraph, mapping: Dict[str, str]) -> bool:
    """Replace placeholders across run boundaries while keeping run formatting."""
    if not mapping or not paragraph.runs:
        return False
    runs = _runs_with_positions(paragraph)
    full_text = "".join(run.text or "" for run in paragraph.runs)
    matches = _find_matches(full_text, mapping)
    if not matches:
        return False

    pieces: List[Tuple[str, Any]] = []
    match_index = 0
    for run, run_start, run_end in runs:
        boundaries = {run_start, run_end}
        for start, end, _value in matches:
            if end <= run_start or start >= run_end:
                continue
            boundaries.add(max(start, run_start))
            boundaries.add(min(end, run_end))
        ordered = sorted(boundaries)
        for left, right in zip(ordered, ordered[1:]):
            if left == right:
                continue
            covering = None
            for start, end, value in matches:
                if start <= left and right <= end:
                    covering = (start, end, value)
                    break
            if covering is not None:
                if left == covering[0]:
                    pieces.append((covering[2], run))
                continue
            original = (run.text or "")[left - run_start : right - run_start]
            if original:
                pieces.append((original, run))
        match_index += 1

    if not pieces:
        for start, end, value in matches:
            if value:
                pieces.append((value, runs[0][0]))
    if not pieces:
        for run, _start, _end in runs:
            try:
                run._r.getparent().remove(run._r)
            except (AttributeError, ValueError):
                pass
        return True

    first_run_element = runs[0][0]._r
    parent = first_run_element.getparent()
    anchor_index = parent.index(first_run_element)
    for run, _start, _end in runs:
        try:
            parent.remove(run._r)
        except ValueError:
            pass
    for text, source_run in pieces:
        new_run = OxmlElement("w:r")
        if source_run._r.rPr is not None:
            new_run.append(copy.deepcopy(source_run._r.rPr))
        for part_index, part in enumerate(str(text).split("\n")):
            if part_index:
                new_run.append(OxmlElement("w:br"))
            if part:
                text_el = OxmlElement("w:t")
                text_el.set(qn("xml:space"), "preserve")
                text_el.text = part
                new_run.append(text_el)
        parent.insert(anchor_index, new_run)
        anchor_index += 1
    return True


def replace_in_container(container, mapping: Dict[str, str]) -> int:
    count = 0
    for paragraph in container.paragraphs:
        if replace_in_paragraph(paragraph, mapping):
            count += 1
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                count += replace_in_container(cell, mapping)
    return count
